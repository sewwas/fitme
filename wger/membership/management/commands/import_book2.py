"""
Fit Me Management Command: Import historical members and subscriptions from Book2.xlsx.
Usage:
    python manage.py import_book2 --dry-run
    python manage.py import_book2 --commit
"""
import re
from datetime import datetime, time
from decimal import Decimal

from django.core.management.base import BaseCommand
from django.contrib.auth.models import User, Group
from django.db import transaction
from django.utils import timezone

import openpyxl

from wger.membership.models import MembershipPlan, MemberProfile, Subscription
from wger.gym_operations_payroll.models import PaymentReceipt


class Command(BaseCommand):
    help = 'Import and sync 78 members from Book2.xlsx into Fit Me database.'

    def add_arguments(self, parser):
        parser.add_argument(
            '--file',
            type=str,
            default='Book2.xlsx',
            help='Path to Excel file (default: Book2.xlsx)'
        )
        parser.add_argument(
            '--dry-run',
            action='store_true',
            help='Simulate the import and display summary without modifying database'
        )
        parser.add_argument(
            '--commit',
            action='store_true',
            help='Write changes to database'
        )

    def handle(self, *args, **options):
        file_path = options['file']
        dry_run = options['dry_run']
        commit = options['commit']

        if not commit and not dry_run:
            self.stdout.write(self.style.WARNING(
                "Neither --commit nor --dry-run specified. Running in DRY-RUN mode for safety.\n"
                "To save changes to the database, re-run with: python manage.py import_book2 --commit\n"
            ))
            dry_run = True

        self.stdout.write(self.style.HTTP_INFO(f"Loading {file_path}..."))
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            sheet = wb.active
        except Exception as e:
            self.stderr.write(self.style.ERROR(f"Failed to open {file_path}: {e}"))
            return

        records = self._parse_records(sheet)
        self.stdout.write(f"Parsed {len(records)} member entries from {file_path}.\n")

        plans_map = self._load_plans()
        member_group, _ = Group.objects.get_or_create(name='member')
        today = timezone.now().date()

        created_users = 0
        updated_users = 0
        created_subs = 0
        created_receipts = 0
        errors = []

        # Print header for dry run / log
        self.stdout.write(f"{'PIN':<6} | {'NAME':<24} | {'PHONE':<12} | {'PLAN':<16} | {'JOIN DATE':<10} | {'VALID UNTIL':<10} | {'PAYMENT'}")
        self.stdout.write("-" * 96)

        try:
            with transaction.atomic():
                for rec in records:
                    pin = rec['pin']
                    name = rec['name']
                    phone = rec['phone']
                    raw_pkg = rec['package']
                    join_date = rec['join_date']
                    pay_date = rec['pay_date']
                    end_date = rec['end_date']
                    dob = rec['dob']
                    email = rec['email']
                    payment_val = rec['payment']
                    section = rec['section']

                    plan = self._resolve_plan(raw_pkg, plans_map)
                    plan_name = plan.name if plan else 'Monthly Package'

                    # Clean phone
                    phone_str = ''
                    if phone:
                        phone_digits = re.sub(r'\D', '', str(phone))
                        if len(phone_digits) == 9:
                            phone_str = f"0{phone_digits}"
                        elif len(phone_digits) >= 10:
                            phone_str = phone_digits
                        else:
                            phone_str = str(phone)

                    # Payment amount (fallback to plan standard price if omitted in Excel)
                    try:
                        amount = Decimal(str(payment_val)) if payment_val and str(payment_val).replace('.', '', 1).isdigit() else (plan.price_monthly if plan else Decimal('0.00'))
                    except Exception:
                        amount = plan.price_monthly if plan else Decimal('0.00')

                    self.stdout.write(
                        f"{str(pin):<6} | {name[:24]:<24} | {phone_str:<12} | {plan_name[:16]:<16} | "
                        f"{str(join_date or '—'):<10} | {str(end_date or '—'):<10} | LKR {amount}"
                    )

                    if not dry_run:
                        # Check existing user by PIN or phone
                        profile = MemberProfile.objects.filter(biometric_pin=str(pin)).select_related('user').first()
                        if not profile and phone_str:
                            profile = MemberProfile.objects.filter(phone=phone_str).select_related('user').first()

                        if profile:
                            user = profile.user
                            updated_users += 1
                            # Update join date if not already set
                            if join_date and not profile.join_date:
                                profile.join_date = join_date
                            if dob and not profile.date_of_birth:
                                profile.date_of_birth = dob
                            if not profile.biometric_pin:
                                profile.biometric_pin = str(pin)
                            profile.save()
                        else:
                            # Generate unique username
                            base_uname = re.sub(r'[^a-zA-Z0-9]', '_', name.lower().strip())
                            base_uname = re.sub(r'_+', '_', base_uname).strip('_') or 'member'
                            username = f"{base_uname}_{pin}"
                            if User.objects.filter(username=username).exists():
                                username = f"{username}_{pin}"

                            user_email = email if (email and '@' in str(email)) else f"{username}@fitmefitness.lk"

                            name_parts = name.split()
                            first_name = name_parts[0]
                            last_name = " ".join(name_parts[1:]) if len(name_parts) > 1 else ""

                            user = User.objects.create(
                                username=username,
                                email=user_email,
                                first_name=first_name,
                                last_name=last_name,
                                is_active=True
                            )
                            user.set_password('Fitme@123')
                            user.save()
                            user.groups.add(member_group)

                            profile = MemberProfile.objects.create(
                                user=user,
                                biometric_pin=str(pin),
                                phone=phone_str,
                                join_date=join_date or today,
                                date_of_birth=dob,
                                gender='male',
                                notes=f"Imported from Book2.xlsx ({'New Member' if section == 'new_members' else 'Renewal'})"
                            )
                            created_users += 1

                        # Resolve dates for subscription
                        sub_start = pay_date or join_date or today
                        sub_end = end_date or (sub_start + timezone.timedelta(days=plan.duration_days if plan else 30))
                        is_active_sub = (sub_end >= today)
                        sub_status = 'active' if is_active_sub else 'expired'

                        sub = Subscription.objects.create(
                            member=user,
                            plan=plan or plans_map['individual'],
                            start_date=sub_start,
                            end_date=sub_end,
                            status=sub_status,
                            payment_method='cash',
                            amount_paid=amount,
                            payment_date=timezone.make_aware(datetime.combine(sub_start, time(9, 0))) if sub_start else timezone.now(),
                            notes=f"Imported from Book2.xlsx (row PIN {pin})"
                        )
                        created_subs += 1

                        # Create PaymentReceipt
                        receipt_no = PaymentReceipt.generate_next_receipt_number()
                        # Calculate annual registration fee component
                        annual_fee = Decimal('0.00')
                        pkg_amount = amount
                        if section == 'new_members' and amount >= Decimal('5000.00') and plan.plan_type != 'student':
                            annual_fee = Decimal('1000.00')
                            pkg_amount = amount - annual_fee

                        pay_datetime = timezone.make_aware(datetime.combine(sub_start, time(9, 0))) if sub_start else timezone.now()

                        PaymentReceipt.objects.create(
                            receipt_number=receipt_no,
                            subscription=sub,
                            member=user,
                            plan_name=plan.name if plan else 'Monthly Package',
                            duration_days=plan.duration_days if plan else 30,
                            validity_start=sub_start,
                            validity_end=sub_end,
                            amount_paid=amount,
                            package_amount=pkg_amount,
                            annual_fee_amount=annual_fee,
                            payment_method='cash',
                            member_pin=str(pin),
                            notes="Imported from Book2.xlsx September 2026 ledger",
                            issued_at=pay_datetime
                        )
                        created_receipts += 1

                if dry_run:
                    # In dry run, rollback any simulated db changes if any
                    transaction.set_rollback(True)

        except Exception as e:
            self.stderr.write(self.style.ERROR(f"\nImport encountered an error: {e}"))
            import traceback
            traceback.print_exc()
            return

        self.stdout.write("-" * 96)
        if dry_run:
            self.stdout.write(self.style.SUCCESS(
                f"\n[DRY RUN COMPLETE]\n"
                f"Records evaluated: {len(records)}\n"
                f"No changes were written to the database.\n"
                f"To execute the actual sync, run:\n"
                f"  python manage.py import_book2 --commit\n"
            ))
        else:
            self.stdout.write(self.style.SUCCESS(
                f"\n[IMPORT COMMITTED SUCCESSFULLY]\n"
                f"Total Records: {len(records)}\n"
                f"Created Member Users: {created_users}\n"
                f"Updated Existing Profiles: {updated_users}\n"
                f"Created Subscriptions: {created_subs}\n"
                f"Generated Payment Receipts: {created_receipts}\n"
            ))

    def _parse_records(self, sheet):
        """Extracts and standardizes records from both sections of Book2.xlsx."""
        records = []

        def parse_date(val):
            if not val:
                return None
            if isinstance(val, datetime):
                return val.date()
            cleaned = str(val).strip().rstrip('.')
            for fmt in ('%d.%m.%Y', '%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y'):
                try:
                    return datetime.strptime(cleaned, fmt).date()
                except ValueError:
                    pass
            return None

        # Section 1: Renewals (Rows 4 to 58)
        for r in range(4, 59):
            pin = sheet.cell(row=r, column=2).value
            name = sheet.cell(row=r, column=3).value
            if pin is not None and name is not None:
                records.append({
                    'row': r,
                    'section': 'renewals',
                    'pin': pin,
                    'name': str(name).strip(),
                    'phone': sheet.cell(row=r, column=4).value,
                    'package': sheet.cell(row=r, column=5).value,
                    'pay_date': parse_date(sheet.cell(row=r, column=6).value),
                    'payment': sheet.cell(row=r, column=7).value,
                    'join_date': parse_date(sheet.cell(row=r, column=8).value),
                    'end_date': parse_date(sheet.cell(row=r, column=9).value),
                    'dob': None,
                    'email': sheet.cell(row=r, column=10).value
                })

        # Section 2: New Members (Rows 66 to 88)
        for r in range(66, 89):
            pin = sheet.cell(row=r, column=2).value
            name = sheet.cell(row=r, column=3).value
            if pin is not None and name is not None:
                records.append({
                    'row': r,
                    'section': 'new_members',
                    'pin': pin,
                    'name': str(name).strip(),
                    'phone': sheet.cell(row=r, column=4).value,
                    'payment': sheet.cell(row=r, column=5).value,
                    'join_date': parse_date(sheet.cell(row=r, column=6).value),
                    'pay_date': parse_date(sheet.cell(row=r, column=6).value),
                    'dob': parse_date(sheet.cell(row=r, column=7).value),
                    'package': sheet.cell(row=r, column=8).value,
                    'end_date': parse_date(sheet.cell(row=r, column=9).value),
                    'email': sheet.cell(row=r, column=10).value
                })

        return records

    def _load_plans(self):
        """Preload plans into a lookup map by plan_type."""
        plans = {}
        for p in MembershipPlan.objects.all():
            plans[p.plan_type] = p
        return plans

    def _resolve_plan(self, raw_pkg, plans_map):
        """Maps freeform package strings to official MembershipPlan objects."""
        if not raw_pkg:
            return plans_map.get('individual')
        s = str(raw_pkg).lower().strip()
        if 'student' in s:
            return plans_map.get('student')
        if 'cuple' in s or 'couple' in s:
            return plans_map.get('couples')
        if 'family' in s or 'famliy' in s:
            return plans_map.get('family_3')
        if 'off_peak' in s or 'day' in s:
            return plans_map.get('off_peak', plans_map.get('individual'))
        return plans_map.get('individual')
